import os
import pdfplumber
import pytesseract
from PIL import Image
from docx import Document
import openpyxl
import re
from dateutil.parser import parse
from io import BytesIO
from tqdm import tqdm
import hashlib
import sqlite3
import cv2
import numpy as np
from concurrent.futures import ThreadPoolExecutor, as_completed

# ---------- CONFIG ----------
pytesseract.pytesseract.tesseract_cmd = r"/usr/bin/tesseract"

SCAN_ROOT = "EpsteinArchiveFolder"
BASE_OUTPUT = "outputlocation"

FULLTEXT_DIR = os.path.join(BASE_OUTPUT, "fulltext")
FACES_DIR = os.path.join(BASE_OUTPUT, "faces")
DB_PATH = os.path.join(BASE_OUTPUT, "results.db")

os.makedirs(FULLTEXT_DIR, exist_ok=True)
os.makedirs(FACES_DIR, exist_ok=True)

MAX_WORKERS = 3
FACE_DETECTION_ENABLED = True

# ---------- KEYWORDS ----------
KEYWORDS = [
    "epstein","jeffrey epstein","jeffery epstein","jeffery",
    "epstein island","little st. james","lsj","little saint james",
    "palm beach","new york mansion","massage","trafficking","minor",
    "victim","flight log","flight logs","pilot",

    "maxwell","ghislaine","gmax","wexner","leslie wexner",
    "ehud barak","jean-luc","brunel","jean-luc brunel",
    "kellen","dubin","glenn dubin","ava dubin",

    "trump","donald trump","melania","ivanka","mar-a-lago",
    "clinton","bill clinton","hillary clinton","prince andrew",
    "kevin spacey","bill gates","dershowitz","alan dershowitz",

    "giuffre","virginia giuffre","virginia roberts",
    "johanna sjoberg","courtney wilde","carolyn",

    "zorro ranch","st. thomas","virgin islands","upper east side",

    "affidavit","testimony","indictment","deposition",
    "lawsuit","sealed","unsealed","fbi","cia","mi6","plea deal",

    "flight","manifest","passport","travel records",

    "bank wire","trust","foundation","blackmail","surveillance",
    "recordings","hidden camera",

    "escort","modeling agency","handler"
]

# ---------- HELPERS ----------

def clean(t):
    return re.sub(r"\s+", " ", t.strip())

def extract_dates(text):
    out = []
    for word in text.split():
        try:
            dt = parse(word, fuzzy=False, ignoretz=True)
            out.append(str(dt.date()))
        except:
            pass
    return sorted(set(out))

def extract_names(text):
    return sorted(set(re.findall(r"\b[A-Z][a-z]+\s[A-Z][a-z]+\b", text)))

def save_fulltext(text, real_file):
    h = hashlib.sha1(real_file.encode("utf-8")).hexdigest()
    out = f"{h}.txt"
    with open(os.path.join(FULLTEXT_DIR, out), "w", encoding="utf-8") as f:
        f.write(text)
    return out

def detect_faces(img, real_file):
    if not FACE_DETECTION_ENABLED or img is None:
        return 0, ""
    try:
        gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        cascade = cv2.CascadeClassifier(
            cv2.data.haarcascades + "haarcascade_frontalface_default.xml")
        faces = cascade.detectMultiScale(gray, 1.1, 5, minSize=(40,40))

        if len(faces) == 0:
            return 0, ""

        h = hashlib.sha1(real_file.encode("utf-8")).hexdigest()
        saved = []

        for idx, (x,y,w,h2) in enumerate(faces):
            face = img[y:y+h2, x:x+w]
            fn = f"{h}_face{idx}.jpg"
            cv2.imwrite(os.path.join(FACES_DIR, fn), face)
            saved.append(fn)

        return len(faces), ", ".join(saved)
    except:
        return 0, ""

# ---------- EXTRACTORS ----------

def load_image(data):
    arr = np.frombuffer(data, np.uint8)
    return cv2.imdecode(arr, cv2.IMREAD_COLOR)

def ocr_image(img):
    try:
        rgb = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
        pil = Image.fromarray(rgb)
        return clean(pytesseract.image_to_string(pil))
    except:
        return ""

def extract_pdf(data):
    out = ""
    with pdfplumber.open(BytesIO(data)) as pdf:
        for p in pdf.pages:
            out += (p.extract_text() or "") + "\n"
    return clean(out)

def extract_docx(data):
    doc = Document(BytesIO(data))
    return clean("\n".join(p.text for p in doc.paragraphs))

def extract_xl(data):
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    out = ""
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out += " ".join(str(v) for v in row if v) + " "
    return clean(out)

def extract_txt(data):
    return clean(data.decode("utf-8", errors="ignore"))

# ---------- DB ----------

def init_db():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            real_file TEXT,
            keywords TEXT,
            names TEXT,
            dates TEXT,
            snippet TEXT,
            fulltext_file TEXT,
            has_faces INTEGER,
            face_files TEXT
        )
    """)
    conn.commit()
    conn.close()

def insert_hit(row):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO results
        (real_file, keywords, names, dates, snippet, fulltext_file, has_faces, face_files)
        VALUES (?,?,?,?,?,?,?,?)
    """, (
        row["real_file"],
        row["keywords"],
        row["names"],
        row["dates"],
        row["snippet"],
        row["fulltext_file"],
        row["has_faces"],
        row["face_files"],
    ))
    conn.commit()
    conn.close()

# ---------- PROCESS ONE FILE ----------

def process(path):
    ext = path.lower()
    try:
        data = open(path, "rb").read()
        real = os.path.abspath(path)

        text = ""

        if ext.endswith((".jpg",".jpeg",".png",".tif",".tiff")):
            img = load_image(data)
            text = ocr_image(img)
            faces, facefiles = detect_faces(img, real)

        elif ext.endswith(".pdf"):
            text = extract_pdf(data)
            faces, facefiles = 0, ""

        elif ext.endswith(".docx"):
            text = extract_docx(data)
            faces, facefiles = 0, ""

        elif ext.endswith((".xlsx",".xls")):
            text = extract_xl(data)
            faces, facefiles = 0, ""

        elif ext.endswith(".txt"):
            text = extract_txt(data)
            faces, facefiles = 0, ""

        else:
            return None

        low = text.lower()

        # ---------- KEYWORD MATCH FILTER ----------
        found = [k for k in KEYWORDS if k in low]
        if not found:
            return None   # ← skip NON-MATCHES

        snippet = text[:2000]
        names = ", ".join(extract_names(text))
        dates = ", ".join(extract_dates(text))
        fulltxt = save_fulltext(text, real)

        return {
            "real_file": real,
            "keywords": ", ".join(found),
            "names": names,
            "dates": dates,
            "snippet": snippet,
            "fulltext_file": fulltxt,
            "has_faces": faces,
            "face_files": facefiles,
        }

    except:
        return None

# ---------- MAIN ----------

def main():
    init_db()

    files = []
    for root, _, fs in os.walk(SCAN_ROOT):
        for f in fs:
            files.append(os.path.join(root, f))

    print(f"Found {len(files)} files\n")
    print("Scanning... (v10 keyword-only mode)\n")

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as ex:
        futures = {ex.submit(process, p): p for p in files}
        for fut in tqdm(as_completed(futures), total=len(futures)):
            hit = fut.result()
            if hit:
                insert_hit(hit)

    print("\nDone.\nDB:", DB_PATH)

if __name__ == "__main__":
    main()
